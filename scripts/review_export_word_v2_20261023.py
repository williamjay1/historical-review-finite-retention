"""Export the revised article sources with manual numbered citations."""
from pathlib import Path
import os, subprocess, re, json, hashlib, zipfile, shutil
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

R = Path(__file__).resolve().parents[1]
M = R / 'manuscript'
O = R / 'results/revision_review_20261023'
T = R / 'temp/review_word_v2_20261023'
PANDOC = Path(os.environ.get('PANDOC', shutil.which('pandoc') or 'pandoc'))
WIDTHS = {
    'ESWA_Manuscript': [[43, 20, 23, 23, 23, 43], [59, 25, 33, 30, 28], [48, 25, 25, 25, 26, 26], [51, 22, 28, 25, 25, 24]],
    'ESWA_Supplementary': [[67, 27, 27, 27, 27], [51, 31, 31, 31, 31], [59, 29, 30, 28, 29], [59, 29, 30, 28, 29], [25, 28, 30, 30, 31, 31]],
}


def border(parent, edge, val='nil', sz=6):
    x = OxmlElement('w:' + edge)
    x.set(qn('w:val'), val)
    x.set(qn('w:sz'), str(sz))
    x.set(qn('w:color'), '000000')
    parent.append(x)


def tables(doc, stem):
    assert len(doc.tables) == len(WIDTHS[stem]), (len(doc.tables), len(WIDTHS[stem]))
    for tb, widths in zip(doc.tables, WIDTHS[stem]):
        assert len(widths) == len(tb.columns) and sum(widths) == 175
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tb.autofit = False
        pr = tb._tbl.tblPr
        old = pr.find(qn('w:tblBorders'))
        if old is not None:
            pr.remove(old)
        bd = OxmlElement('w:tblBorders')
        for edge in ['top', 'bottom']:
            border(bd, edge, 'single', 8)
        for edge in ['left', 'right', 'insideH', 'insideV']:
            border(bd, edge)
        pr.append(bd)
        mar = OxmlElement('w:tblCellMar')
        for side in ['top', 'bottom', 'left', 'right']:
            x = OxmlElement('w:' + side)
            x.set(qn('w:w'), '65' if side in ['top', 'bottom'] else '80')
            x.set(qn('w:type'), 'dxa')
            mar.append(x)
        pr.append(mar)
        for col, w in zip(tb.columns, widths):
            col.width = Mm(w)
        for i, row in enumerate(tb.rows):
            rp = row._tr.get_or_add_trPr()
            rp.append(OxmlElement('w:cantSplit'))
            if i == 0:
                rp.append(OxmlElement('w:tblHeader'))
            for j, (cell, w) in enumerate(zip(row.cells, widths)):
                cell.width = Mm(w)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr()
                for sh in list(cp.findall(qn('w:shd'))):
                    cp.remove(sh)
                cb = OxmlElement('w:tcBorders')
                for ed in ['top', 'bottom', 'left', 'right']:
                    outer = (i == 0 and ed == 'top') or (i == len(tb.rows) - 1 and ed == 'bottom')
                    header = i == 0 and ed == 'bottom'
                    border(cb, ed, 'single' if outer or header else 'nil', 8 if outer else 6)
                cp.append(cb)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.keep_with_next = i < len(tb.rows) - 1
                    for run in p.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(9 if stem == 'ESWA_Manuscript' else 8.5)
                        run.bold = i == 0


def main():
    T.mkdir(parents=True, exist_ok=True)
    ref = Document()
    sec = ref.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(19)
    sec.left_margin = sec.right_margin = Mm(17.5)
    norm = ref.styles['Normal']
    norm.font.name = 'Times New Roman'
    norm.font.size = Pt(11)
    norm.paragraph_format.space_after = Pt(6)
    norm.paragraph_format.widow_control = True
    for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'Heading 5', 'Heading 6']:
        st = ref.styles[name]
        st.font.name = 'Arial'
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.keep_with_next = True
    ref.styles['Title'].font.size = Pt(16)
    ref.styles['Heading 1'].font.size = Pt(13)
    ref.styles['Heading 2'].font.size = Pt(11.5)
    vb = ref.styles.add_style('Verbatim Char', WD_STYLE_TYPE.CHARACTER)
    vb._element.set(qn('w:styleId'), 'VerbatimChar')
    vb.font.name = 'Consolas'
    vb.font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    fp._p.append(fld)
    ref.save(T / 'reference.docx')
    receipts = []
    for stem, src_stem, nfig in [
        ('ESWA_Manuscript', 'ESWA_Manuscript', 2),
        ('ESWA_Supplementary', 'ESWA_Supplementary', 0),
    ]:
        src = M / (src_stem + '.md')
        dst = M / (stem + '.docx')
        cmd = [str(PANDOC), str(src), '--from=markdown+tex_math_dollars+tex_math_single_backslash',
               '--to=docx', '--standalone', '--reference-doc=' + str(T / 'reference.docx'),
               '--resource-path=' + str(M), '-o', str(dst)]
        p = subprocess.run(cmd, cwd=M, capture_output=True, text=True, encoding='utf-8', errors='replace')
        (T / (stem + '_pandoc.log')).write_text(p.stdout + p.stderr, encoding='utf-8')
        assert p.returncode == 0 and not p.stderr.strip(), p.stderr
        d = Document(dst)
        tables(d, stem)
        for p in d.paragraphs:
            if p.text.startswith('Table '):
                p.paragraph_format.keep_with_next = True
            if p._p.xpath('.//w:drawing'):
                p.paragraph_format.keep_with_next = True
            if p.text.startswith('Fig. '):
                p.paragraph_format.space_after = Pt(7)
                for run in p.runs:
                    run.font.size = Pt(10)
        for shape in d.inline_shapes:
            ratio = shape.height / shape.width
            shape.width = Mm(175)
            shape.height = int(shape.width * ratio)
        d.save(dst)
        with zipfile.ZipFile(dst) as z:
            assert z.testzip() is None
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                  'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
                  'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            tree = etree.fromstring(z.read('word/document.xml'))
            text = ' '.join(tree.xpath('//w:t/text()', namespaces=ns))
            nf = len(tree.xpath('//a:blip', namespaces=ns))
            nt = len(tree.xpath('//w:tbl', namespaces=ns))
            nm = len(tree.xpath('//m:oMath', namespaces=ns))
            assert nf == nfig, (nf, nfig)
            assert nt == len(WIDTHS[stem]), (nt, len(WIDTHS[stem]))
            assert nm > 0 or stem == 'ESWA_Supplementary'
            assert not re.search(r'\[@\w+|[A-Z]:[\\/]|MLWork|TODO|\{\{', text)
        receipts.append(dict(file=str(dst), bytes=dst.stat().st_size,
                             sha256=hashlib.sha256(dst.read_bytes()).hexdigest(),
                             source_sha256=hashlib.sha256(src.read_bytes()).hexdigest(),
                             native_tables=nt, inline_figures=nf, native_equations=nm,
                             citations_resolved=True, render_status='pending native Word inspection'))
    (O / 'WORD_DELIVERY_CHECK.json').write_text(json.dumps(receipts, indent=2), encoding='utf-8')
    print(json.dumps(receipts, indent=2))


if __name__ == '__main__':
    main()
